#!/usr/bin/env python3
"""동일한 답안·자료를 사용한 사람·LLM 채점자의 항목별 일치도를 계산한다."""
from __future__ import annotations

import argparse
import itertools
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import common  # noqa: E402
import rubric  # noqa: E402
import score  # noqa: E402
import grading


def item_max(rb: dict) -> dict[str, float]:
    return {it["id"]: it["points"] for _, issue in rubric.issues(rb) for it in rubric.items(issue)}


def from_judgment(j: dict) -> dict[str, float | None]:
    return {r["id"]: r["awarded"] for v in j["issues"].values() for r in v["items"]}


def from_sheet(path: Path, rb: dict) -> tuple[dict[str, float | None], dict[str, float | None]]:
    """사람 채점표 → (항목 점수, 문서별 재량점수). 빈 칸은 None."""
    ws = load_workbook(path, data_only=True)["채점"]
    items, disc, seen = {}, {}, set()
    expected = {item["id"]: (doc, issue, item) for doc, issue, item in grading.entries(rb)}
    for row in ws.iter_rows(min_row=2, values_only=True):
        iid, val = row[0], row[8]
        if not isinstance(iid, str) or row[3] is None or iid == "합계(참고)":
            continue
        if iid not in expected or iid in seen:
            raise ValueError(f"{path.name}: 알 수 없거나 중복된 항목 {iid}")
        seen.add(iid)
        doc, issue, item = expected[iid]
        if val is not None and not rubric.valid_score(val, item):
            raise ValueError(f"{path.name}/{iid}: 허용되지 않은 사람 점수 {val!r}")
        if issue is None:
            disc[doc] = val
        else:
            items[iid] = val
    if seen != set(expected):
        raise ValueError(f"{path.name}: 평가항목 누락")
    return items, disc


def as_judgment(items: dict, disc: dict, rb: dict) -> dict:
    """사람 점수를 공통 합산 함수의 입력 구조로 변환한다."""
    return {"issues": {issue["id"]: {"items": [{"id": it["id"], "awarded": items.get(it["id"])}
                                               for it in rubric.items(issue)]}
                       for _, issue in rubric.issues(rb)},
            "discretion": {d["id"]: {"id": d["discretion"]["id"], "awarded": disc.get(d["id"])}
                           for d in rb["documents"] if d.get("discretion")}}


def alpha_interval(pairs: list[tuple[float, float]]) -> float | None:
    """두 채점자, 결측 없음 — Krippendorff α (구간 척도)."""
    if len(pairs) < 2:
        return None
    d_o = sum((a - b) ** 2 for a, b in pairs) / len(pairs)
    vals = [v for p in pairs for v in p]
    n = len(vals)
    mean = sum(vals) / n
    d_e = 2 * sum((x - mean) ** 2 for x in vals) / (n - 1)
    return None if d_e == 0 else 1 - d_o / d_e


def compare(a: dict, b: dict, mx: dict) -> dict:
    ids = [i for i in mx if a.get(i) is not None and b.get(i) is not None]
    for iid in ids:
        if mx[iid] <= 0 or any(type(r[iid]) not in (int, float) or not 0 <= r[iid] <= mx[iid] for r in (a, b)):
            raise ValueError(f"일치도 점수 범위 오류: {iid}")
    if not ids:
        return {"n": 0}
    same = sum(1 for i in ids if a[i] == b[i])
    gap = sum(abs(a[i] - b[i]) / mx[i] for i in ids) / len(ids)
    return {"n": len(ids), "same": same / len(ids), "gap": gap,
            "alpha": alpha_interval([(a[i] / mx[i], b[i] / mx[i]) for i in ids])}


def fmt(c: dict) -> str:
    if not c["n"]:
        return "비교할 항목 없음"
    al = "—" if c["alpha"] is None else f"{c['alpha']:.2f}"
    return f"항목 {c['n']}개 · 같은 점수 {c['same']:.0%} · 평균 차이(배점 대비) {c['gap']:.0%} · α {al}"


def main() -> None:
    ap = argparse.ArgumentParser(description="채점자 간 일치도")
    ap.add_argument("-t", "--task", required=True)
    ap.add_argument("--judges", nargs="*", default=[])
    ap.add_argument("--human", action="store_true", help="results/human/<과제>/ 채점표도 비교")
    ap.add_argument("--human-raters", nargs="+", help="사람별 하위 폴더에서 독립 채점표를 읽음")
    ap.add_argument("-m", "--models", nargs="*", help="이 답안만 비교 (기본: 전부, 모범답안은 reference)")
    ap.add_argument("--results", type=Path, default=common.RESULTS)
    common.add_data_argument(ap)
    args = ap.parse_args()
    common.unique_names(args.judges, "채점 모델")
    common.unique_names(args.human_raters or [], "사람 채점자")
    if len(args.judges) + len(args.human_raters or []) + bool(args.human) < 2:
        ap.error("비교할 채점자를 두 명 이상 지정해야 함")

    rb = common.load_rubric(args.task, args.data_root)
    mx = item_max(rb)
    raters: dict[str, dict[str, dict]] = {}  # 채점자 → 답안 → 항목 점수
    totals: dict[str, dict[str, float | None]] = {}
    failures = 0
    answer_hashes = {}
    for jd in args.judges:
        for f in sorted((args.results / "judgments" / common.slug(jd) / args.task).glob("*.json")):
            j = common.read_json(f)
            if args.models and j["model"] not in args.models:
                continue
            errors = score.source_errors(j, rb, args.data_root, args.results) + score.judgment_errors(j, rb)
            if errors:
                failures += 1
                print(f"[agreement] 미완료/이전 결과 제외: {f} — {'; '.join(errors)}", file=sys.stderr)
                continue
            if j["model"] in answer_hashes and answer_hashes[j["model"]] != j["answer_sha"]:
                sys.exit("[agreement] 같은 이름에 서로 다른 답안이 연결됨")
            answer_hashes[j["model"]] = j["answer_sha"]
            raters.setdefault(jd, {})[j["model"]] = from_judgment(j)
            totals.setdefault(jd, {})[j["model"]] = score.score(j, rb)["content"]
    human_raters = ([None] if args.human else []) + (args.human_raters or [])
    context = grading.load_context(rb, args.data_root) if rb["meta"].get("reference_guide") else None
    current_inputs = {"rubric_sha": common.sha(common.rubric_file(args.task, args.data_root).read_text(encoding="utf-8")),
                      "record_sha": common.sha(common.load_record(args.task, numbered=True, root=args.data_root)),
                      "source_hashes": grading.source_hashes(context["sources"]) if context else {}}
    for rater in human_raters:
        name = "사람" if rater is None else f"사람:{rater}"
        hdir = args.results / "human" / args.task
        if rater is not None:
            hdir = hdir / common.slug(rater)
        saved_key = common.read_json(hdir / "key.json")
        if saved_key.get("inputs") != current_inputs:
            sys.exit("[agreement] 사람 채점표의 루브릭·기록·규칙·기준답안 버전이 다름")
        for label, model in sorted(saved_key["ids"].items()):
            if args.models and model not in args.models:
                continue
            path = hdir / f"{label}.xlsx"
            try:
                digest = saved_key["answer_sha"][label]
                if model in answer_hashes and digest != answer_hashes[model]:
                    raise ValueError("사람과 LLM 또는 사람끼리 다른 답안을 채점함")
                items, disc = from_sheet(path, rb)
                total = score.score(as_judgment(items, disc, rb), rb)
                if not total["complete"]:
                    raise ValueError("사람 채점 미완료: " + "; ".join(total["errors"]))
                answer_hashes[model] = digest
                raters.setdefault(name, {})[model] = items
                totals.setdefault(name, {})[model] = total["content"]
            except (OSError, ValueError, KeyError) as error:
                failures += 1
                print(f"[agreement] {name}/{label}: {error}", file=sys.stderr)

    if len(raters) < 2:
        sys.exit("[agreement] 현재 기준으로 비교할 채점 결과 없음")
    names = list(raters)
    expected_raters = set(args.judges) | {"사람" if r is None else f"사람:{r}" for r in human_raters}
    if expected_raters - set(raters):
        failures += 1
        print("[agreement] 요청한 채점자 중 완료된 결과가 없는 대상이 있음", file=sys.stderr)
    answers = sorted({m for r in raters.values() for m in r} | set(args.models or []))
    content_max = score.score(as_judgment({}, {}, rb), rb)["content_max"]
    print(f"## {args.task} — 답안별 내용 점수 (재량 제외, /{content_max:g})\n")
    print("| 답안 | " + " | ".join(names) + " |")
    print("|---|" + "---|" * len(names))
    for m in answers:
        cells = (totals.get(n, {}).get(m) for n in names)
        print(f"| {m} | " + " | ".join("—" if c is None else f"{c:g}" for c in cells) + " |")

    print("\n## 항목 단위 일치도\n")
    for x, y in itertools.combinations(names, 2):
        common_answers = [m for m in answers if m in raters[x] and m in raters[y]]
        if len(common_answers) != len(answers):
            failures += 1
            print(f"[agreement] {x} ↔ {y}: 일부 답안의 채점 결과 누락", file=sys.stderr)
        pooled_a, pooled_b = {}, {}
        for m in common_answers:
            for i in mx:
                pooled_a[f"{m}/{i}"], pooled_b[f"{m}/{i}"] = raters[x][m].get(i), raters[y][m].get(i)
        pooled_mx = {f"{m}/{i}": mx[i] for m in common_answers for i in mx}
        print(f"- **{x} ↔ {y}** (답안 {len(common_answers)}개): {fmt(compare(pooled_a, pooled_b, pooled_mx))}")
        for m in common_answers:
            print(f"  - {m}: {fmt(compare(raters[x][m], raters[y][m], mx))}")

    return 1 if failures else 0


if __name__ == "__main__":
    common.cli(main)
