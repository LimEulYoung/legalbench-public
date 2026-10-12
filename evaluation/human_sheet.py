#!/usr/bin/env python3
"""익명 답안과 채점표·기록·규칙·기준답안을 사람 채점자별로 생성한다."""
from __future__ import annotations

import argparse
import math
import random
import re
import string
import sys
import unicodedata
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
import common  # noqa: E402
import rubric  # noqa: E402
import grading

KIND = {"disposition": "결론", "bonus": "가점"}
COLS = [("항목", 9), ("구분", 14), ("유형", 6), ("배점", 6), ("채점 기준", 60), ("점수별 조건", 45),
        ("추가 채점 지침", 24), ("기록 근거", 24), ("점수", 7), ("메모", 30)]
SCORE_COL, MAX_COL = "I", "D"  # 점수 입력 칸, 배점 칸 (COLS 순서와 맞출 것)
INPUT = PatternFill("solid", fgColor="FFF2CC")
ISSUE = PatternFill("solid", fgColor="D9D9D9")
WRAP = Alignment(wrap_text=True, vertical="top")

GUIDE = """\
채점 방법
- 「채점」 시트의 노란 칸(점수)에 항목의 점수표에 있는 값 하나를 넣는다. 메모는 선택.
- 답안 원문은 「답안」 시트나 같은 이름의 .md 파일에 있다. 어느 모델의 답안인지는 가려 두었다(key.json 은 열지 말 것).
- 문제·작성요령·주의사항과 기록 전문은 「기록」 시트에서 확인한다.
- 채점 모델의 점수는 보지 않고 채점한다. 답안에 실제로 쓰인 내용을 평가하고 원용·중복 감점 방지 규칙을 함께 적용한다.
- 가점 항목은 해당 내용이 있을 때만 준다. 문서별 상한은 합산할 때 코드로 적용하므로 신경 쓰지 않아도 된다.
- 재량점수도 해당 행의 기준과 점수별 조건에 따라 준다. 개별 내용 점수를 중복 지급하지 않는다.

채점 규칙
"""


def items_rows(rb: dict):
    """(종류, 값들) — 종류는 'issue'(쟁점 머리줄) · 'item' · 'discretion'."""
    for doc in rb["documents"]:
        for issue in doc["issues"]:
            yield "issue", (f"{issue['id']} {issue['title']} ({issue['points']:g}점)",)
            for sec in issue["sections"]:
                for it in sec["items"]:
                    guidance = it.get("accept", "")
                    if guidance and it.get("accept_label"):
                        guidance = f"{it['accept_label']}: {guidance}"
                    guidance = "\n\n".join(filter(None, (guidance, rubric.explanation_text(it))))
                    yield "item", (it["id"], sec["title"], KIND.get(it.get("type"), "일반"), it["points"],
                                   it["criterion"], rubric.scoring_text(it), guidance,
                                   it.get("record", ""))
        if doc.get("discretion"):
            it = doc["discretion"]
            yield "discretion", (it["id"], doc["title"], "재량", it["points"], it["criterion"],
                                 rubric.scoring_text(it), "", "")


def add_reading_guide(wb, rb, context):
    """기준답안 발췌와 항목별 해설을 세로로 읽을 수 있는 시트."""
    sheet = wb.create_sheet("채점안내", 0)
    sheet.column_dimensions["A"].width = 120
    sheet.freeze_panes = "A2"

    def write(text, bold=False):
        for line in text.splitlines():
            sheet.append([line])
            cell = sheet.cell(sheet.max_row, 1)
            cell.alignment = WRAP
            cell.font = Font(bold=bold)
            width = sum(2 if unicodedata.east_asian_width(char) in "WF" else 1 for char in line)
            sheet.row_dimensions[sheet.max_row].height = min(409, 18 * max(1, math.ceil(width / 110)))

    write(f"{rb['meta']['exam']} {rb['meta']['subject']} 채점안내", True)
    write(f"v{rb['meta'].get('version', '0.1')} · {rb['meta']['status']}")
    write("쟁점별 기준답안을 읽고 항목의 채점 기준·점수 조건·해설을 확인한다. "
          "실제 답안의 연결된 문단도 함께 읽는다. 점수와 메모는 「채점」 시트에 입력한다.")
    excerpts = rubric.reference_excerpts(rb, context["reference"] if context else None)
    positions = rubric.sheet_positions(rb)
    for doc, issue in rubric.issues(rb):
        write(f"\n{positions[issue['id']]} {issue['title']} ({issue['points']:g}점)", True)
        for excerpt in excerpts.get(issue["id"], []):
            write(f"기준답안 — {excerpt['heading']}", True)
            write(re.sub(r"(?m)^#{2,6}\s+", "", excerpt["text"]))
        for item in rubric.items(issue):
            write(f"\n{positions[item['id']]} {item['criterion']} ({item['points']:g}점)", True)
            if item.get("explanation"):
                write("채점 해설: " + item["explanation"])
            write(rubric.scoring_text(item))
            if item.get("accept"):
                write(f"{item.get('accept_label', '허용 답안')}: {item['accept']}")


def build_sheet(path: Path, answer: str, rb: dict, label: str, record=None, context=None) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "채점"
    ws.append([c for c, _ in COLS])
    for i, (_, w) in enumerate(COLS, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        ws.cell(1, i).font = Font(bold=True)
    ws.freeze_panes = "A2"
    all_items = [it for _, issue in rubric.issues(rb) for it in rubric.items(issue)]
    all_items += [doc["discretion"] for doc in rb["documents"] if doc.get("discretion")]
    allowed = {it["id"]: ",".join(f"{v:g}" for v in sorted(it["scoring"])) for it in all_items}
    first = None
    for kind, vals in items_rows(rb):
        if kind == "issue":
            ws.append([vals[0]])
            r = ws.max_row
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(COLS))
            ws.cell(r, 1).fill, ws.cell(r, 1).font = ISSUE, Font(bold=True)
            continue
        ws.append(list(vals) + [None, None])
        r = ws.max_row
        first = first or r
        for c in range(1, len(COLS) + 1):
            ws.cell(r, c).alignment = WRAP
        ws[f"{SCORE_COL}{r}"].fill = INPUT
        dv = DataValidation(type="list", formula1=f'"{allowed[vals[0]]}"', allow_blank=True,
                            showErrorMessage=True, errorTitle="점수 오류", error="이 항목의 점수표에서 선택하세요.")
        ws.add_data_validation(dv)
        dv.add(ws[f"{SCORE_COL}{r}"])
    last = ws.max_row
    ws.append([])
    ws.append(["합계(참고)", None, None, f"=SUM({MAX_COL}{first}:{MAX_COL}{last})", None, None, None, None,
               f"=SUM({SCORE_COL}{first}:{SCORE_COL}{last})"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)

    wa = wb.create_sheet("답안")
    wa.column_dimensions["A"].width = 120
    wa.append([f"답안 {label}"])
    for para in answer.splitlines():
        wa.append([para])
        wa.cell(wa.max_row, 1).alignment = WRAP

    wg = wb.create_sheet("안내")
    wg.column_dimensions["A"].width = 120
    rules = context["rules"] if context else "\n".join(rb["meta"]["grading_rules"])
    for line in (GUIDE + rules).splitlines():
        wg.append([line])
        wg.cell(wg.max_row, 1).alignment = WRAP
    for title, text in (("기록", record), ("기준답안", context["reference"] if context else None)):
        if text is not None:
            sheet = wb.create_sheet(title)
            sheet.column_dimensions["A"].width = 120
            for line in text.splitlines():
                sheet.append([line])
                sheet.cell(sheet.max_row, 1).alignment = WRAP
    if any(q.get("reference_sections") or any(it.get("explanation") for it in rubric.items(q))
           for _, q in rubric.issues(rb)):
        add_reading_guide(wb, rb, context)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    ap = argparse.ArgumentParser(description="사람 채점용 블라인드 채점표 생성")
    ap.add_argument("-t", "--task", required=True)
    ap.add_argument("-m", "--models", nargs="*", help="대상 모델 (기본: 답안이 있는 모든 모델)")
    ap.add_argument("--results", type=Path, default=common.RESULTS)
    common.add_data_argument(ap)
    ap.add_argument("--rater", help="사람 채점자 식별자; 사람별 폴더에 독립 채점표 생성")
    ap.add_argument("--reference", type=Path, help="함께 섞을 모범답안 파일 (이름은 reference)")
    ap.add_argument("--force", action="store_true", help="이미 있는 채점표도 다시 만듦(입력한 점수가 지워짐)")
    args = ap.parse_args()

    rb = common.load_rubric(args.task, args.data_root)
    record = common.load_record(args.task, numbered=True, root=args.data_root)
    context = grading.load_context(rb, args.data_root) if rb["meta"].get("reference_guide") else None
    errs = rubric.validate(rb)
    if errs:
        sys.exit("[human] 루브릭 오류: " + "; ".join(errs))
    adir = args.results / "answers" / args.task
    texts = {}
    for path in adir.glob("*.json"):
        if path.name.endswith(".failed.json"):
            continue
        data = common.read_json(path)
        model = data.get("model")
        if not isinstance(model, str) or path.name != common.slug(model) + ".json" or data.get("task", args.task) != args.task:
            sys.exit(f"[human] 답안 모델·과제 식별자 불일치: {path}")
        texts[model] = common.answer_text(path)
        if not texts[model].strip():
            sys.exit(f"[human] 빈 답안: {path}")
    models = args.models or sorted(texts)
    missing = [m for m in models if m not in texts]
    if missing:
        sys.exit(f"[human] 답안 없음: {', '.join(missing)} ({adir})")
    if args.reference:
        texts["reference"] = re.sub(r"<!--.*?-->", "", args.reference.read_text(encoding="utf-8"), flags=re.S).strip()
        models = models + ["reference"]

    out = args.results / "human" / args.task
    if args.rater:
        common.unique_names([args.rater], "채점자")
        out = out / common.slug(args.rater)
    key_path = out / "key.json"
    key = common.read_json(key_path) if key_path.exists() else {"task": args.task, "ids": {}}
    rubric_sha = common.sha(common.rubric_file(args.task, args.data_root).read_text(encoding="utf-8"))
    inputs = {"rubric_sha": rubric_sha, "record_sha": common.sha(record),
              "source_hashes": grading.source_hashes(context["sources"]) if context else {}}
    if key_path.exists() and key.get("rubric_sha") != rubric_sha and not args.force:
        sys.exit("[human] 이전 루브릭의 채점표가 있음. 별도 results 폴더를 사용하거나 --force로 다시 생성할 것")
    if key_path.exists() and not args.force:
        if key.get("inputs") != inputs or any(key.get("answer_sha", {}).get(label) != common.sha(texts[model])
                for label, model in key["ids"].items() if model in models):
            sys.exit("[human] 기록·기준답안·규칙·답안이 변경됨. 새 결과 폴더를 사용할 것")
    if args.force and key.get("inputs") not in (None, inputs) and not set(key["ids"].values()).issubset(models):
        sys.exit("[human] 자료가 바뀐 묶음은 전체 답안을 재생성하거나 새 결과 폴더를 사용할 것")
    have = set(key["ids"].values())
    new = [m for m in models if m not in have]
    random.SystemRandom().shuffle(new)  # 모델 순서가 ID 로 드러나지 않게
    free = [c for c in string.ascii_uppercase if c not in key["ids"]]
    if len(new) > len(free):
        sys.exit("[human] 한 묶음은 최대 26개 답안. 별도 결과 폴더로 나누어 생성할 것")
    for m in new:
        key["ids"][free.pop(0)] = m
    key["rubric_sha"] = rubric_sha
    key["inputs"] = inputs
    key.setdefault("answer_sha", {}).update({label: common.sha(texts[model])
        for label, model in key["ids"].items() if model in models})
    common.write_json(key_path, key)

    for label, model in sorted(key["ids"].items()):
        if model not in models:
            continue
        xlsx = out / f"{label}.xlsx"
        if xlsx.exists() and not args.force:
            print(f"[human] 건너뜀(있음): {xlsx}", file=sys.stderr)
            continue
        answer = texts[model]
        (out / f"{label}.md").write_text(f"# 답안 {label} — {rb['meta']['exam']} {rb['meta']['subject']}\n\n"
                                         f"{answer.strip()}\n", encoding="utf-8")
        build_sheet(xlsx, answer, rb, label, record, context)
        print(f"[human] {xlsx} ({len(answer):,}자)", file=sys.stderr)
    print(f"[human] 채점표 {len(key['ids'])}개 · 모델 대응표: {key_path} (채점 전에는 열지 말 것)", file=sys.stderr)


if __name__ == "__main__":
    common.cli(main)
