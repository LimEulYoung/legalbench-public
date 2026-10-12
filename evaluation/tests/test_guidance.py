"""공개 가상 자료로 사람·AI 안내의 일치와 기준답안 연결을 검증한다."""
import copy
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import common
import grading
import human_sheet
import judge
import rubric

FIXTURE = Path(__file__).resolve().parents[1] / "examples"
REFERENCE = "# 기준답안\n\n## 운영 안내\n\n### 시간\n\n운영 시간을 안내합니다.\n\n### 휴관\n\n월요일은 쉽니다.\n"


class GuidanceTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.rb = common.load_rubric("r1-public", FIXTURE)
        self.rb["meta"].update(status="검토 초안", guidelines="rules.md", reference_guide="reference.md")
        (self.root / "rules.md").write_text("## 판정 규칙\n답안의 실제 내용만 채점한다.\n", encoding="utf-8")
        self.reference = self.root / "reference.md"
        self.reference.write_text(REFERENCE, encoding="utf-8")
        self.issue = self.rb["documents"][0]["issues"][0]
        self.issue["reference_sections"] = ["운영 안내 / 시간"]
        self.item = rubric.items(self.issue)[0]
        self.item.update(explanation="시작과 종료 시간을 모두 확인한다.",
                         note="내부 비교 메모는 제공하지 않는다.")
        self.answer = "제가 제출한 답안입니다."

    def test_shared_guidance_and_reference_in_markdown_workbook_and_prompt(self):
        self.assertEqual(rubric.validate(self.rb), [])
        context = grading.load_context(self.rb, self.root)
        md = rubric.render(self.rb, reference_text=context["reference"])
        prompt = grading.messages(self.rb, self.answer, "가상 기록", context, judge.format_item)[1]["content"]
        path = self.root / "sheet.xlsx"
        human_sheet.build_sheet(path, self.answer, self.rb, "A", record="가상 기록", context=context)
        wb = load_workbook(path)
        self.addCleanup(wb.close)
        self.assertEqual(wb.sheetnames[0], "채점안내")
        guide = "\n".join(str(row[0] or "") for row in wb["채점안내"].values)
        for text in (self.item["explanation"], "운영 시간을 안내합니다."):
            for output in (md, guide, prompt):
                self.assertIn(text, output)
        for output in (md, guide, prompt):
            self.assertNotIn(self.item["note"], output)
            self.assertNotIn("점수 예시", output)
        self.assertNotIn("월요일은 쉽니다.", md)
        self.assertEqual(prompt.count(REFERENCE), 1)
        ws = wb["채점"]
        rows = {row[0].value: row for row in ws if row[0].value in {"I-1", "I-2"}}
        self.assertEqual(set(rows), {"I-1", "I-2"})
        self.assertEqual(ws["I1"].value, "점수")
        for row in rows.values():
            self.assertIsNone(row[8].value)
            self.assertIsNone(row[9].value)
        self.assertIn(self.item["explanation"], rows["I-1"][6].value)
        self.assertEqual(rows["I-1"][6].value, rubric.explanation_text(self.item))

    def test_reference_edits_reach_next_render_and_hash(self):
        first = grading.load_context(self.rb, self.root)
        self.reference.write_text(REFERENCE.replace("운영 시간을 안내합니다.", "수정된 운영 안내입니다."), encoding="utf-8")
        second = grading.load_context(self.rb, self.root)
        self.assertNotEqual(first["sources"], second["sources"])
        self.assertIn("수정된 운영 안내입니다.", rubric.render(self.rb, reference_text=second["reference"]))

    def test_missing_empty_and_ambiguous_reference_sections_fail(self):
        cases = [REFERENCE.replace("### 시간", "### 다른 제목"),
                 REFERENCE.replace("운영 시간을 안내합니다.", ""),
                 REFERENCE + "\n### 시간\n\n중복된 제목입니다.\n",
                 REFERENCE.replace("운영 시간을 안내합니다.", "") + "\n### 시간\n\n중복입니다.\n"]
        for text in cases:
            with self.subTest(text=text), self.assertRaises(ValueError):
                rubric.reference_excerpts(self.rb, text)
        with self.assertRaises(ValueError):
            rubric.reference_excerpts(self.rb, None)
        self.issue["reference_sections"].append("운영 안내 / 시간")
        self.assertTrue(rubric.validate(self.rb))

    def test_unusable_explanation_is_rejected(self):
        for explanation in ("", " ", None, 1, []):
            item = copy.deepcopy(self.item)
            item["explanation"] = explanation
            with self.subTest(explanation=explanation):
                self.assertTrue(rubric.validate_item(item))

    def test_reference_or_explanation_cannot_be_used_as_candidate_evidence(self):
        for quote in (self.item["explanation"], "운영 시간을 안내합니다."):
            row, errors = judge.check_item({"id": "I-1", "awarded": 1, "quotes": [quote],
                                           "record_lines": [], "reason": "근거를 인용함"},
                                          self.item, self.answer, "가상 기록")
            self.assertTrue(any("답안 원문에 없는 인용" in error for error in errors))
            self.assertIsNone(row["awarded"])

    def test_existing_rubrics_do_not_require_new_fields(self):
        self.issue.pop("reference_sections")
        self.item.pop("explanation")
        self.assertEqual(rubric.validate(self.rb), [])
        self.assertEqual(rubric.reference_excerpts(self.rb, None), {})
        self.assertNotIn("점수 예시", rubric.render(self.rb))
        path = self.root / "old-sheet.xlsx"
        human_sheet.build_sheet(path, self.answer, self.rb, "A")
        wb = load_workbook(path)
        self.addCleanup(wb.close)
        self.assertNotIn("채점안내", wb.sheetnames)


if __name__ == "__main__":
    unittest.main()
