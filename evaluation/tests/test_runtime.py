"""공개 가상 자료와 로컬 HTTP 서버로 실행·복구·재집계 경계를 검증한다."""
from __future__ import annotations

import copy
import io
import itertools
import json
import os
import shutil
import subprocess
import sys
import tempfile
import threading
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import agreement
import common
import grading
import generation_report
import human_sheet
import judge
import llm
import score

EVAL = Path(__file__).resolve().parents[1]
FIXTURE = EVAL / 'examples'
ANSWER = (FIXTURE / 'answer.md').read_text(encoding='utf-8')
TASK = 'r1-public'


def event(data):
    text = data if isinstance(data, str) else json.dumps(data, ensure_ascii=False)
    return ('data: ' + text + '\n\n').encode('utf-8')


def judgment():
    return {'items': [
        {'id': 'I-1', 'awarded': 1, 'quotes': ['도서관은 오전 9시부터 오후 6시까지 운영합니다.'],
         'record_lines': [6], 'reason': '두 운영 시각을 썼다.', 'alternative': '',
         'review_required': False, 'review_reason': ''},
        {'id': 'I-2', 'awarded': 1, 'quotes': ['월요일은 휴관입니다.'],
         'record_lines': [], 'reason': '휴관일을 썼다.', 'alternative': '',
         'review_required': False, 'review_reason': ''}]}


class TransportTest(unittest.TestCase):
    def test_sse_unicode_usage_and_metadata(self):
        stream = event({'id': 'request-1', 'model': 'served', 'choices': [{'delta': {'content': '한글'}}]})
        stream += event({'choices': [{'delta': {}, 'finish_reason': 'stop'}]})
        stream += event({'choices': [], 'usage': {'prompt_tokens': 20, 'completion_tokens': 3}})
        stream += event('[DONE]')
        result = llm._read_stream(io.BytesIO(stream))
        summary = llm.summarize(result, {'served': (0.01, 0.02)}, 'requested')
        self.assertEqual(summary['content'], '한글')
        self.assertEqual(summary['response_id'], 'request-1')
        self.assertAlmostEqual(summary['cost_usd'], 0.26)
        self.assertEqual(summary['cost_source'], 'estimated')

    def test_damaged_or_interrupted_stream_preserves_partial_response(self):
        prefix = event({'choices': [{'delta': {'content': 'partial'}}]})
        for tail in (b'data: {broken\n\n', b'data: \xff\n\n', b'', event('[DONE]')):
            with self.subTest(tail=tail), self.assertRaises(llm.LLMError) as caught:
                llm._read_stream(io.BytesIO(prefix + tail))
            self.assertEqual(caught.exception.response['choices'][0]['message']['content'], 'partial')

    def test_unknown_cost_is_null_and_reported_zero_is_zero(self):
        response = {'choices': [{'message': {'content': 'ok'}, 'finish_reason': 'stop'}]}
        for usage in ({}, {'prompt_tokens': 10}, {'completion_tokens': 10}, {'cost': -1}):
            self.assertIsNone(llm.summarize(dict(response, usage=usage), {'m': (1, 2)}, 'm')['cost_usd'])
        self.assertEqual(llm.summarize(dict(response, usage={'cost': 0}), {}, 'm')['cost_usd'], 0)

    def test_api_stream_error_keeps_details_and_partial_response_without_key(self):
        stream = event({'choices': [{'delta': {'content': 'partial'}}]})
        stream += event({'error': {'code': 'upstream_timeout', 'message': 'request test-key failed'}})
        with patch.object(llm.urllib.request, 'urlopen', return_value=io.BytesIO(stream)) as request:
            with self.assertRaisesRegex(llm.LLMError, 'upstream_timeout') as caught:
                llm.chat('test-key', [], 'm', stream=True)
            request.assert_called_once()
        self.assertNotIn('test-key', str(caught.exception))
        self.assertIn('[redacted]', str(caught.exception))
        self.assertEqual(caught.exception.response['choices'][0]['message']['content'], 'partial')

    def test_partial_transport_failure_and_http_error_are_single_requests(self):
        import urllib.error
        with patch.object(llm.urllib.request, 'urlopen', side_effect=urllib.error.HTTPError(
                'http://localhost', 429, 'busy', {}, io.BytesIO(b'busy'))) as request:
            with self.assertRaisesRegex(llm.LLMError, '429'):
                llm.chat('test-key', [], 'm', stream=True)
            request.assert_called_once()

    def test_http_failure_preserves_status_request_id_and_redacted_raw_body(self):
        error = llm.urllib.error.HTTPError('http://localhost', 429, 'busy',
            {'x-request-id': 'http-request-1'}, io.BytesIO(b'busy test-key'))
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'response.txt'
            with patch.object(llm.urllib.request, 'urlopen', side_effect=error):
                with self.assertRaises(llm.LLMError) as caught:
                    llm.chat('test-key', [], 'm', stream=True, raw_response_path=path)
            self.assertEqual(path.read_bytes(), b'busy [redacted]')
            self.assertEqual(caught.exception.kind, 'http_error')
            self.assertEqual(caught.exception.http_status, 429)
            self.assertEqual(caught.exception.response['gateway_request_id'], 'http-request-1')

    def test_chat_completion_token_field_and_raw_capture(self):
        stream = event({'choices': [{'delta': {'content': '답안'}, 'finish_reason': 'stop'}]}) + event('[DONE]')
        with tempfile.TemporaryDirectory() as folder:
            raw = Path(folder) / 'response.txt'
            with patch.object(llm.urllib.request, 'urlopen', return_value=io.BytesIO(stream)) as request:
                llm.chat('test-key', [], 'm', stream=True, max_tokens=131072,
                         token_limit_field='max_completion_tokens', raw_response_path=raw)
            self.assertEqual(raw.read_bytes(), stream)
            payload = json.loads(request.call_args.args[0].data)
            self.assertEqual(payload['max_completion_tokens'], 131072)
            self.assertNotIn('max_tokens', payload)

    def test_structured_refusals_survive_chat_and_responses_normalization(self):
        chat_stream = event({'choices': [{'delta': {'refusal': '거절'}, 'finish_reason': 'stop'}]}) + event('[DONE]')
        response_stream = event({'type': 'response.refusal.delta', 'delta': '거절'})
        response_stream += event({'type': 'response.completed', 'response': {'status': 'completed'}})
        responses = [llm._read_stream(io.BytesIO(chat_stream)),
                     llm._read_responses_stream(io.BytesIO(response_stream)),
                     llm._normalize_response({'status': 'completed', 'output': [{'type': 'message',
                         'content': [{'type': 'refusal', 'refusal': '거절'}]}]})]
        for response in responses:
            self.assertEqual(llm.summarize(response, {}, 'm')['refusal'], '거절')

    def test_responses_request_normalizes_text_reasoning_usage_and_status(self):
        stream = event({'type': 'response.reasoning_summary_text.delta', 'delta': '검토'})
        stream += event({'type': 'response.output_text.delta', 'delta': '답안'})
        stream += event({'type': 'response.completed', 'response': {
            'id': 'resp-1', 'model': 'm', 'status': 'completed', 'output': [],
            'usage': {'input_tokens': 100, 'output_tokens': 20,
                      'output_tokens_details': {'reasoning_tokens': 15}},
            'estimated_cost': {'currency': 'unit', 'total': 1}}})
        with patch.object(llm.urllib.request, 'urlopen', return_value=io.BytesIO(stream)) as request:
            response = llm.chat('test-key', [{'role': 'user', 'content': '문제'}], 'm', stream=True,
                                api_format='responses', max_tokens=200, reasoning_effort='max')
        sent = request.call_args.args[0]
        payload = json.loads(sent.data)
        self.assertTrue(sent.full_url.endswith('/responses'))
        self.assertEqual(payload['input'], [{'role': 'user', 'content': '문제'}])
        self.assertEqual(payload['max_output_tokens'], 200)
        self.assertEqual(payload['reasoning'], {'effort': 'max'})
        self.assertFalse({'messages', 'max_tokens', 'reasoning_effort', 'stream_options'} & payload.keys())
        summary = llm.summarize(response, {}, 'm')
        self.assertEqual((summary['content'], summary['reasoning'], summary['finish_reason']), ('답안', '검토', 'stop'))
        self.assertEqual(summary['usage']['completion_tokens_details']['reasoning_tokens'], 15)
        self.assertIsNone(summary['cost_usd'])
        self.assertEqual(summary['gateway_cost']['currency'], 'unit')

    def test_responses_completed_event_can_report_incomplete_output(self):
        final = {'status': 'incomplete', 'incomplete_details': {'reason': 'max_output_tokens'},
                 'output': [{'type': 'reasoning', 'summary': [{'type': 'summary_text', 'text': '검토'}]}]}
        response = llm._read_responses_stream(io.BytesIO(event({'type': 'response.completed', 'response': final})))
        summary = llm.summarize(response, {}, 'm')
        self.assertEqual(summary['finish_reason'], 'length')
        self.assertEqual(summary['reasoning'], '검토')
        self.assertEqual(llm._normalize_response(final), response)

    def test_responses_failed_and_interrupted_streams_preserve_partial_text(self):
        prefix = event({'type': 'response.reasoning_text.delta', 'delta': '검토'})
        prefix += event({'type': 'response.output_text.delta', 'delta': '일부 답안'})
        tails = (b'', b'data: {broken\n\n', event({'type': 'response.failed', 'response': {
            'status': 'failed', 'error': {'code': 'upstream_timeout'}}}),
            event({'type': 'response.completed', 'response': {'status': 'completed', 'output': ['invalid']}}))
        for tail in tails:
            with self.subTest(tail=tail), self.assertRaises(llm.LLMError) as caught:
                llm._read_responses_stream(io.BytesIO(prefix + tail))
            summary = llm.summarize(caught.exception.response, {}, 'm')
            self.assertEqual(summary['content'], '일부 답안')
            self.assertEqual(summary['reasoning'], '검토')

    def test_non_usd_price_cache_is_not_used_as_dollars(self):
        usage = [{'unit': 'token_input', 'amount': 1}, {'unit': 'token_output', 'amount': 2}]
        models = [{'id': currency or 'legacy', 'pricing': {'currency': currency, 'usage': usage}}
                  for currency in ('unit', 'USD', None)]
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'prices.json'
            common.write_json(path, models)
            self.assertEqual(llm.load_prices(path), {'USD': (1, 2), 'legacy': (1, 2)})


class RuntimeTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.data = self.root / 'materials'
        shutil.copytree(FIXTURE, self.data)
        self.results = self.root / 'results'
        self.requests = []
        self.mode = 'normal'
        owner = self

        class Handler(BaseHTTPRequestHandler):
            def log_message(self, *args):
                pass

            def do_POST(self):
                payload = json.loads(self.rfile.read(int(self.headers['Content-Length'])))
                owner.requests.append(payload)
                mode = owner.mode.pop(0) if isinstance(owner.mode, list) else owner.mode
                messages = payload.get('messages') or payload['input']
                is_judge = '<rubric>' in messages[-1]['content']
                content = json.dumps(judgment(), ensure_ascii=False) if is_judge else ANSWER
                if mode == 'invalid-json' and is_judge:
                    content = 'invalid JSON'
                if mode == 'empty':
                    content = ' \n '
                self.send_response(200)
                self.send_header('Content-Type', 'text/event-stream')
                self.send_header('x-letsur-request-id', 'gateway-local-1')
                self.end_headers()
                chunks = [event({'id': 'local-1', 'model': payload['model'],
                                 'choices': [{'index': 0, 'delta': {'content': content}}]})]
                usage = {'prompt_tokens': 100, 'completion_tokens': 20, 'cost': 0.02,
                         'completion_tokens_details': {'reasoning_tokens': 0}}
                if mode == 'reasoning':
                    chunks.append(event({'choices': [{'delta': {'reasoning_content': 'Test reasoning.'}}]}))
                    usage['completion_tokens_details']['reasoning_tokens'] = 5
                if mode == 'broken-stream':
                    chunks.append(b'data: {broken\n\n')
                else:
                    finish = mode if mode in {'length', 'content_filter', 'tool_calls'} else 'stop'
                    chunks += [event({'choices': [{'delta': {}, 'finish_reason': finish}],
                                      'usage': usage}), event('[DONE]')]
                if self.path.endswith('/responses'):
                    chunks = [event({'type': 'response.output_text.delta', 'delta': content}),
                              event({'type': 'response.completed', 'response': {
                                  'id': 'local-1', 'model': payload['model'], 'status': 'completed',
                                  'usage': {'input_tokens': 100, 'output_tokens': 20, 'cost': 0.02}}})]
                for chunk in chunks:
                    self.wfile.write(chunk)
                self.wfile.flush()

        self.server = ThreadingHTTPServer(('127.0.0.1', 0), Handler)
        thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        thread.start()
        self.addCleanup(self.server.server_close)
        self.addCleanup(self.server.shutdown)
        self.url = f'http://127.0.0.1:{self.server.server_port}/v1'
        self.env = dict(os.environ, LEGALBENCH_TEST_KEY='local-only-dummy-key')

    def cli(self, script, *extra, status=0, results=None, max_tokens=4096):
        argv = [sys.executable, '-B', str(EVAL / script), '--results', str(results or self.results)]
        if script in ('judge.py', 'generate.py'):
            argv += ['-t', TASK, '--data-root', str(self.data), '--base-url', self.url,
                     '--api-key-env', 'LEGALBENCH_TEST_KEY', '--seed', '7']
            if max_tokens is not None:
                argv += ['--max-tokens', str(max_tokens)]
        command = subprocess.run(argv + list(extra), capture_output=True, text=True, env=self.env, timeout=30)
        self.assertEqual(command.returncode, status, command.stdout + command.stderr)
        return command

    def candidate(self):
        common.write_json(common.answer_file(self.results, TASK, 'candidate'),
                          {'task': TASK, 'model': 'candidate', 'content': ANSWER})

    def path(self):
        return common.judgment_file(self.results, TASK, 'candidate', 'local-judge')

    def test_full_pipeline_and_replay_after_move_without_original_inputs(self):
        self.cli('generate.py', '-m', 'candidate')
        self.cli('judge.py', '-m', 'candidate', '--judge', 'local-judge')
        self.assertEqual(len(self.requests), 2)
        self.assertEqual(self.requests[0]['seed'], 7)
        result = common.read_json(self.path())
        self.assertEqual(result['totals']['total'], 2)
        manifests = [common.read_json(p) for p in (self.results / 'runs').glob('*.json')]
        self.assertEqual({m['stage'] for m in manifests}, {'judge', 'generate'})
        self.assertTrue(all(m['complete'] for m in manifests))
        self.assertTrue(manifests[0]['environment']['code_sha256'])
        self.assertNotIn('local-only-dummy-key', json.dumps(manifests))
        moved = self.root / 'moved-results'
        self.results.rename(moved)
        shutil.rmtree(self.data)
        csv = moved / 'scores.csv'
        self.cli('score.py', '--judge', 'local-judge', '--csv', str(csv), results=moved)
        self.assertTrue(csv.is_file())
        self.cli('score.py', '--judge', 'local-judge', '--source', 'current', '--data-root', str(self.data),
                 results=moved, status=1)

    def test_missing_answer_exits_failure_and_remains_in_experiment(self):
        self.candidate()
        self.cli('judge.py', '-m', 'candidate', 'missing', '--judge', 'local-judge', status=1)
        run = common.read_json(next((self.results / 'runs').glob('*-judge.json')))
        self.assertEqual(run['counts'], {'completed': 1, 'failed': 1})
        self.assertEqual(len(self.requests), 1)
        csv = self.results / 'invalid.csv'
        command = self.cli('score.py', '--judge', 'local-judge', '--csv', str(csv), status=1)
        self.assertIn('예정된 결과 누락', command.stderr)
        self.assertFalse(csv.exists())

    def test_responses_generation_and_judging_preserve_api_configuration(self):
        self.cli('generate.py', '-m', 'candidate', '--api-format', 'responses', '--reasoning-effort', 'max')
        self.cli('judge.py', '-m', 'candidate', '--judge', 'local-judge', '--api-format', 'responses')
        answer = common.read_json(common.answer_file(self.results, TASK, 'candidate'))
        result = common.read_json(self.path())
        self.assertEqual(answer['params']['api_format'], 'responses')
        self.assertEqual(result['params']['api_format'], 'responses')
        self.assertEqual(result['totals']['total'], 2)
        self.assertEqual(self.requests[0]['reasoning'], {'effort': 'max'})
        self.assertTrue(all('max_output_tokens' in request for request in self.requests))
        self.cli('generate.py', '-m', 'candidate', '--reasoning-effort', 'max', status=1)
        self.assertEqual(len(self.requests), 2)

    def test_completion_token_field_reaches_generation_and_judging(self):
        self.cli('generate.py', '-m', 'candidate', '--token-limit-field', 'max_completion_tokens')
        self.cli('judge.py', '-m', 'candidate', '--judge', 'local-judge',
                 '--token-limit-field', 'max_completion_tokens')
        self.assertEqual(len(self.requests), 2)
        for request in self.requests:
            self.assertEqual(request['max_completion_tokens'], 4096)
            self.assertNotIn('max_tokens', request)

    def test_astra_profile_reaches_api_and_saved_settings(self):
        self.candidate()
        self.cli('judge.py', '-m', 'candidate', '--judge', 'gpt-6-astra',
                 '--profile', 'gpt-6-astra', max_tokens=None)
        self.assertEqual(len(self.requests), 1)
        request = self.requests[0]
        self.assertEqual(request['model'], 'gpt-6-astra')
        self.assertEqual(request['max_output_tokens'], 128000)
        self.assertEqual(request['reasoning'], {'effort': 'high'})
        self.assertNotIn('max_tokens', request)
        path = common.judgment_file(self.results, TASK, 'candidate', 'gpt-6-astra')
        result = common.read_json(path)
        run = common.read_json(next((self.results / 'runs').glob('*-judge.json')))
        for settings in (result['params'], run['settings']):
            self.assertEqual(settings['api_format'], 'responses')
            self.assertEqual(settings['max_tokens'], 128000)
            self.assertEqual(settings['reasoning_effort'], 'high')
        self.assertTrue(result['totals']['complete'])

    def test_astra_profile_rejects_conflicting_options_before_request(self):
        for options in (('--max-tokens', '65536'), ('--reasoning-effort', 'low'),
                        ('--api-format', 'chat'), ('--judge', 'other-judge')):
            with self.subTest(options=options):
                command = self.cli('judge.py', '-m', 'candidate', '--judge', 'gpt-6-astra',
                                   '--profile', 'gpt-6-astra', *options, max_tokens=None, status=2)
                self.assertIn('--profile과', command.stderr)
        self.assertEqual(self.requests, [])
        self.assertFalse(self.results.exists())

    def test_candidate_identity_mismatch_is_rejected_before_request(self):
        self.candidate()
        path = common.answer_file(self.results, TASK, 'candidate')
        common.write_json(path, {'task': TASK, 'model': 'different-model', 'content': ANSWER})
        self.cli('judge.py', '-m', 'candidate', '--judge', 'local-judge', status=1)
        self.assertEqual(self.requests, [])

    def test_invalid_response_and_retry_preserve_history_and_cost(self):
        self.candidate()
        self.mode = 'invalid-json'
        self.cli('judge.py', '-m', 'candidate', '--judge', 'local-judge', status=1)
        self.assertIsNone(common.read_json(self.path())['totals']['total'])
        self.mode = 'normal'
        self.cli('judge.py', '-m', 'candidate', '--judge', 'local-judge', '--retry-failed')
        result = common.read_json(self.path())
        self.assertEqual(result['totals']['total'], 2)
        self.assertEqual(len(result['previous_attempts']), 1)
        self.assertEqual(result['cost_usd'], 0.04)

    def test_broken_stream_saves_partial_body_and_unknown_cost(self):
        self.candidate()
        self.mode = 'broken-stream'
        self.cli('judge.py', '-m', 'candidate', '--judge', 'local-judge', status=1)
        result = common.read_json(self.path())
        self.assertIsNone(result['cost_usd'])
        self.assertTrue(result['call']['content'])
        self.assertIsNone(result['totals']['total'])
        self.assertEqual(len(self.requests), 1)

    def test_generation_refuses_old_settings_and_retains_failed_attempts(self):
        self.mode = ['length', 'empty', 'normal']
        self.cli('generate.py', '-m', 'candidate')
        path = common.answer_file(self.results, TASK, 'candidate')
        result = common.read_json(path)
        self.assertEqual(len(result['previous_attempts']), 2)
        self.assertEqual(result['selected_attempt'], 3)
        self.assertAlmostEqual(result['total_cost_usd'], 0.06)
        self.assertEqual(result['total_usage']['completion_tokens'], 60)
        self.assertEqual(self.requests, [self.requests[0]] * 3)
        raw = self.results / result['raw_response_file']
        self.assertTrue(raw.is_file())
        before = path.read_bytes()
        self.cli('generate.py', '-m', 'candidate', '--temperature', '0.7', status=1)
        self.assertEqual(path.read_bytes(), before)
        self.cli('generate.py', '-m', 'candidate', '--retry-failed')
        self.assertEqual(path.read_bytes(), before)
        self.assertEqual(len(self.requests), 3)
        report = generation_report.build_report(self.results, [TASK, 'r2-public'], ['candidate'])
        self.assertEqual(report['models']['candidate']['first_attempt_completion_rate'], 0)
        self.assertEqual(report['models']['candidate']['completion_rate'], 0.5)
        self.assertEqual(report['models']['candidate']['missing_tasks'], ['r2-public'])

    def test_generation_exhaustion_cannot_be_reset_by_resume_or_force(self):
        self.mode = 'length'
        self.cli('generate.py', '-m', 'candidate', status=1)
        self.cli('generate.py', '-m', 'candidate', '--retry-failed', status=1)
        self.cli('generate.py', '-m', 'candidate', '--force', status=2)
        self.assertEqual(len(self.requests), 3)
        report = generation_report.build_report(self.results)
        self.assertEqual(report['models']['candidate']['attempt_outcomes'], {'length': 3})
        self.assertEqual(report['models']['candidate']['completed'], 0)

    def test_generation_does_not_retry_transport_errors_refusals_or_unknown_finish(self):
        for mode in ('broken-stream', 'content_filter', 'tool_calls'):
            with self.subTest(mode=mode):
                self.mode = mode
                root = self.root / mode
                self.cli('generate.py', '-m', 'candidate', status=1, results=root)
                self.cli('generate.py', '-m', 'candidate', '--retry-failed', status=1, results=root)
                report = generation_report.build_report(root)
                self.assertEqual(report['models']['candidate']['attempts'], 1)
        self.assertEqual(len(self.requests), 3)

    def test_extra_body_is_sent_saved_and_checked_on_resume(self):
        path = self.root / 'request-options.json'
        options = {'thinking': {'type': 'enabled'}, 'reasoning': {'effort': 'max', 'exclude': False}}
        common.write_json(path, options)
        self.cli('generate.py', '-m', 'candidate', '--reasoning-effort', 'max', '--extra-body', str(path))
        answer = common.read_json(common.answer_file(self.results, TASK, 'candidate'))
        self.assertEqual(answer['params']['extra_body'], options)
        self.assertEqual(answer['gateway_request_id'], 'gateway-local-1')
        self.cli('judge.py', '-m', 'candidate', '--judge', 'local-judge', '--extra-body', str(path))
        self.assertEqual(common.read_json(self.path())['params']['extra_body'], options)
        for request in self.requests:
            self.assertEqual({k: request[k] for k in options}, options)
        common.write_json(path, {'thinking': {'type': 'disabled'}})
        self.cli('generate.py', '-m', 'candidate', '--reasoning-effort', 'max', '--extra-body', str(path), status=1)
        self.assertEqual(len(self.requests), 2)

    def test_invalid_extra_body_never_calls_api(self):
        path = self.root / 'request-options.json'
        for body in ([], {'model': 'other'}, {'messages': []}, {'max_tokens': 1}, {'max_output_tokens': 1},
                     {'max_completion_tokens': 1}, {'n': 2},
                     {'api_key': 'not-a-real-key'}, {'reasoning': None},
                     {'reasoning': {'effort': 'low'}}, {'reasoning': {'effort': None}}):
            with self.subTest(body=body):
                common.write_json(path, body)
                self.cli('generate.py', '-m', 'candidate', '--reasoning-effort', 'max',
                         '--extra-body', str(path), status=2)
        path.unlink()
        self.cli('generate.py', '-m', 'candidate', '--extra-body', str(path), status=2)
        self.assertEqual(self.requests, [])

    def test_required_reasoning_failure_is_retained_and_cannot_be_exported(self):
        self.cli('generate.py', '-m', 'candidate', '--require-reasoning', status=1)
        path = common.answer_file(self.results, TASK, 'candidate')
        self.assertFalse(path.exists())
        failed = common.read_json(path.with_name(path.stem + '.failed.json'))
        self.assertEqual(failed['finish_reason'], 'stop')
        self.assertTrue(failed['content'])
        self.assertEqual(failed['reasoning_evidence'],
                         {'observed': False, 'text_characters': 0, 'reported_tokens': 0})
        self.cli('human_sheet.py', '-t', TASK, '-m', 'candidate', '--rater', 'reviewer',
                 '--data-root', str(self.data), status=1)
        self.assertFalse((self.results / 'human').exists())
        self.mode = 'reasoning'
        self.cli('generate.py', '-m', 'candidate', '--require-reasoning', '--retry-failed', status=1)
        self.assertEqual(len(self.requests), 1)
        self.assertEqual(failed['selected_attempt'], 1)
        report = generation_report.build_report(self.results)['models']['candidate']
        self.assertEqual(report['completed'], 1)
        self.assertEqual(report['ready_for_grading'], 0)
        other = self.root / 'verified-settings'
        self.cli('generate.py', '-m', 'candidate', '--require-reasoning', results=other)
        answer = common.read_json(common.answer_file(other, TASK, 'candidate'))
        self.assertTrue(answer['reasoning_evidence']['observed'])
        self.assertEqual(answer['reasoning_evidence']['reported_tokens'], 5)
        self.assertEqual(len(answer['previous_attempts']), 0)
        self.assertEqual(len(self.requests), 2)

    def test_dry_run_and_invalid_arguments_never_call_api(self):
        self.cli('generate.py', '-m', 'candidate', '--dry-run')
        self.cli('judge.py', '--answer-file', str(self.data / 'answer.md'), '--judge', 'local-judge',
                 '--dry-run', '--dump', str(self.root / 'preview'))
        self.assertFalse(self.results.exists())
        for options in (('--max-tokens', '0'), ('--temperature', 'nan'), ('--repeat', '-1'),
                        ('--base-url', 'https://user:secret@example.com/v1')):
            self.cli('judge.py', '--answer-file', str(self.data / 'answer.md'), '--judge', 'local-judge',
                     *options, status=2)
        self.assertEqual(self.requests, [])

    def test_snapshot_tampering_and_mixed_settings_block_scores(self):
        self.candidate()
        self.cli('judge.py', '-m', 'candidate', '--judge', 'local-judge')
        original = common.read_json(self.path())
        bad = copy.deepcopy(original)
        bad['input_snapshot']['rubric'] += '\nchanged'
        common.write_json(self.path(), bad)
        self.cli('score.py', '--judge', 'local-judge', status=1)
        common.write_json(self.path(), original)
        other = copy.deepcopy(original)
        other.update(model='other', params=dict(original['params'], temperature=0.9))
        common.write_json(common.judgment_file(self.results, TASK, 'other', 'local-judge'), other)
        command = self.cli('score.py', '--judge', 'local-judge', status=1)
        self.assertIn('실행 설정', command.stderr)

    def test_human_packages_match_context_and_compare_independent_raters(self):
        from openpyxl import load_workbook
        self.candidate()
        self.cli('judge.py', '-m', 'candidate', '--judge', 'local-judge')
        for name in ('alice', 'bob'):
            self.cli('human_sheet.py', '-t', TASK, '-m', 'candidate', '--rater', name, '--data-root', str(self.data))
            path = next((self.results / 'human' / TASK / name).glob('*.xlsx'))
            workbook = load_workbook(path)
            self.assertTrue({'기록', '기준답안', '안내', '답안'}.issubset(workbook.sheetnames))
            for row in workbook['채점'].iter_rows(min_row=2):
                if row[0].value in {'I-1', 'I-2'}:
                    row[8].value = 1
            workbook.save(path)
        result = self.cli('agreement.py', '-t', TASK, '--judges', 'local-judge', '--human-raters', 'alice', 'bob',
                          '--data-root', str(self.data))
        self.assertIn('같은 점수 100%', result.stdout)


class HumanValidationTest(unittest.TestCase):
    def test_item_validation_rejects_missing_duplicate_unknown_and_bad_scores(self):
        rb = common.load_rubric(TASK, FIXTURE)
        valid = judgment()
        variants = [dict(items=valid['items'][:-1]), dict(items=valid['items'] + [valid['items'][0]])]
        for change in ({'id': 'unknown'}, {'awarded': 99}, {'awarded': True}, {'quotes': ['없는 문장']},
                       {'review_required': True, 'awarded': None, 'review_reason': '사람 판단 필요'}):
            changed = copy.deepcopy(valid)
            changed['items'][0].update(change)
            variants.append(changed)
        record = common.load_record(TASK, numbered=True, root=FIXTURE)
        for data in variants:
            with self.subTest(data=data):
                result = grading.check_response(data, rb, ANSWER, record, judge.check_item)
                self.assertIsNone(score.score(result, rb)['total'])

    def test_invalid_duplicate_and_missing_human_items_are_rejected(self):
        from openpyxl import load_workbook
        rb = common.load_rubric(TASK, FIXTURE)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'sheet.xlsx'
            for variant in ('out-of-range', 'boolean', 'unknown', 'duplicate', 'missing'):
                human_sheet.build_sheet(path, ANSWER, rb, 'A')
                book = load_workbook(path)
                sheet = book['채점']
                rows = [r for r in sheet.iter_rows() if r[0].value in {'I-1', 'I-2'}]
                for row in rows:
                    row[8].value = 1
                if variant == 'out-of-range': rows[0][8].value = 99
                if variant == 'boolean': rows[0][8].value = True
                if variant == 'unknown': rows[0][0].value = 'unknown'
                if variant == 'duplicate': sheet.append([c.value for c in rows[0]])
                if variant == 'missing': sheet.delete_rows(rows[0][0].row)
                book.save(path)
                with self.subTest(variant=variant), self.assertRaises(ValueError):
                    agreement.from_sheet(path, rb)

    def test_interval_alpha_matches_direct_pair_formula(self):
        pairs = [(0, 1), (0.5, 0.5), (1, 0), (0, 0)]
        values = [value for pair in pairs for value in pair]
        observed = sum((a - b) ** 2 for a, b in pairs) / len(pairs)
        expected = sum((a - b) ** 2 for a, b in itertools.permutations(values, 2)) / (len(values) * (len(values) - 1))
        self.assertAlmostEqual(agreement.alpha_interval(pairs), 1 - observed / expected)
        with self.assertRaises(ValueError):
            agreement.compare({'x': 99}, {'x': 1}, {'x': 1})


if __name__ == '__main__':
    unittest.main()
